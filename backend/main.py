import json
import os
from io import BytesIO

from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from openai import OpenAI
from pydantic import BaseModel

load_dotenv()

app = FastAPI(title="CampusX Backend")

api_key = os.getenv("OPENAI_API_KEY")

if not api_key:
    raise RuntimeError("OPENAI_API_KEY is not configured.")

client = OpenAI(api_key=api_key)


# ============================================================
# REQUEST MODELS
# ============================================================


class ExamWarriorRequest(BaseModel):
    subject: str
    unit: str
    topic: str
    difficulty: str


class SafetyAIRequest(BaseModel):
    report_type: str
    description: str
    latitude: float
    longitude: float
    recent_reports: list[dict] = []


class AttendanceAIRequest(BaseModel):
    student_name: str
    present_count: int
    absent_count: int
    total_classes: int
    recent_records: list[dict] = []


class TeacherAttendanceAIRequest(BaseModel):
    class_name: str
    total_students: int
    students: list[dict]
    recent_summary: list[dict] = []


# ============================================================
# BASIC ROUTES
# ============================================================


@app.get("/")
def root():
    return {
        "message": "CampusX Backend is running 🚀"
    }


@app.get("/health")
def health():
    return {
        "status": "ok"
    }


# ============================================================
# SKILLBRIDGE AI
# ============================================================


@app.post("/skillbridge/analyze")
async def skillbridge_analyze(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
):
    if not resume.filename:
        raise HTTPException(
            status_code=400,
            detail="No resume file selected.",
        )

    filename = resume.filename.lower()

    if not filename.endswith((".pdf", ".docx")):
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX resumes are supported.",
        )

    if not job_description.strip():
        raise HTTPException(
            status_code=400,
            detail="Job description is required.",
        )

    try:
        file_bytes = await resume.read()

        if not file_bytes:
            raise HTTPException(
                status_code=400,
                detail="Uploaded resume is empty.",
            )

        # Maximum file size: 10 MB
        if len(file_bytes) > 10 * 1024 * 1024:
            raise HTTPException(
                status_code=400,
                detail="Resume file is too large. Maximum size is 10 MB.",
            )

        # ----------------------------------------------------
        # PDF TEXT EXTRACTION
        # ----------------------------------------------------

        if filename.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(file_bytes))

            text_parts = []

            for page in reader.pages:
                page_text = page.extract_text() or ""

                if page_text.strip():
                    text_parts.append(page_text)

            resume_text = "\n".join(text_parts).strip()

        # ----------------------------------------------------
        # DOCX TEXT EXTRACTION
        # ----------------------------------------------------

        else:
            from docx import Document

            document = Document(BytesIO(file_bytes))

            text_parts = []

            # Normal paragraphs
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()

                if text:
                    text_parts.append(text)

            # Tables
            for table in document.tables:
                for row in table.rows:
                    cells = []

                    for cell in row.cells:
                        cell_text = cell.text.strip()

                        if cell_text:
                            cells.append(cell_text)

                    if cells:
                        text_parts.append(" | ".join(cells))

            resume_text = "\n".join(text_parts).strip()

        # ----------------------------------------------------
        # CHECK EXTRACTED TEXT
        # ----------------------------------------------------

        if not resume_text:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Could not extract text from this resume. "
                    "If this is a scanned/image-only PDF, please "
                    "use a text-based PDF or DOCX."
                ),
            )

        # Limit AI input size
        resume_text = resume_text[:30000]
        job_description = job_description.strip()[:20000]

        # ----------------------------------------------------
        # SKILLBRIDGE AI PROMPT
        # ----------------------------------------------------

        prompt = f"""
You are SkillBridge AI, a career assistant for a college student.

Analyze the student's ACTUAL resume against the provided job
description.

Do not invent skills, experience, projects, certifications,
education, or achievements that are not present in the resume.

RESUME:

{resume_text}

JOB DESCRIPTION:

{job_description}

Calculate an estimated skill match from 0 to 100.

Return ONLY valid JSON in exactly this structure:

{{
  "match_percentage": 0,
  "matched_skills": [],
  "missing_skills": [],
  "roadmap": [
    {{
      "week": 1,
      "title": "",
      "tasks": []
    }},
    {{
      "week": 2,
      "title": "",
      "tasks": []
    }},
    {{
      "week": 3,
      "title": "",
      "tasks": []
    }},
    {{
      "week": 4,
      "title": "",
      "tasks": []
    }}
  ]
}}

Rules:

- match_percentage must be a number between 0 and 100.
- matched_skills must contain only skills actually present
  in the resume and relevant to the job.
- missing_skills must contain skills required or strongly
  preferred by the job but missing from the resume.
- matched_skills and missing_skills must be arrays of strings.
- roadmap must contain exactly 4 weeks.
- Each week must contain a title and an array of practical tasks.
- The roadmap should focus on learning the missing skills.
- Keep the roadmap realistic for a college student.
- Do not claim that the student has skills that are not
  supported by the resume.
- Do not add Markdown.
- Do not add text before or after the JSON.
"""

        response = client.responses.create(
            model="gpt-5.6-luna",
            input=prompt,
        )

        # ----------------------------------------------------
        # CLEAN AI OUTPUT
        # ----------------------------------------------------

        raw_output = response.output_text.strip()

        if raw_output.startswith("```json"):
            raw_output = raw_output[len("```json"):].strip()

        elif raw_output.startswith("```"):
            raw_output = raw_output[len("```"):].strip()

        if raw_output.endswith("```"):
            raw_output = raw_output[:-3].strip()

        start = raw_output.find("{")
        end = raw_output.rfind("}")

        if start != -1 and end != -1 and end > start:
            raw_output = raw_output[start:end + 1]

        result = json.loads(raw_output)

        # ----------------------------------------------------
        # VALIDATE RESULT
        # ----------------------------------------------------

        match_percentage = result.get(
            "match_percentage",
            0,
        )

        if not isinstance(match_percentage, (int, float)):
            match_percentage = 0

        match_percentage = max(
            0,
            min(100, match_percentage),
        )

        matched_skills = result.get(
            "matched_skills",
            [],
        )

        missing_skills = result.get(
            "missing_skills",
            [],
        )

        roadmap = result.get(
            "roadmap",
            [],
        )

        if not isinstance(matched_skills, list):
            matched_skills = []

        if not isinstance(missing_skills, list):
            missing_skills = []

        if not isinstance(roadmap, list):
            roadmap = []

        return {
            "match_percentage": match_percentage,
            "matched_skills": matched_skills,
            "missing_skills": missing_skills,
            "roadmap": roadmap,
            "resume_filename": resume.filename,
        }

    except HTTPException:
        raise

    except json.JSONDecodeError:
        raise HTTPException(
            status_code=500,
            detail="AI returned invalid JSON.",
        )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"SkillBridge AI error: {str(e)}",
        )


# ============================================================
# EXAMWARRIOR AI
# ============================================================


@app.post("/examwarrior/generate")
def examwarrior_generate(request: ExamWarriorRequest):

    prompt = f"""
You are an AI question generator for a B.Tech college
exam preparation app.

Generate ONE useful exam-practice question.

Subject: {request.subject}
Unit: {request.unit}
Topic: {request.topic}
Difficulty: {request.difficulty}

Requirements:

- The question must be relevant to the subject, unit and topic.
- Match the requested difficulty.
- Make it suitable for a B.Tech student.
- Prefer numerical, conceptual, derivation, or university-exam
  style questions depending on the topic.
- Do not provide the answer.
- Do not add explanation.
- Generate exactly ONE question.

Return ONLY this JSON object:

{{
  "question": "Your generated question here"
}}

The "question" field must contain a non-empty string.

Do not use Markdown code fences.
Do not add text before or after the JSON.
"""

    try:
        response = client.responses.create(
            model="gpt-5.6-luna",
            input=prompt,
        )

        raw_output = response.output_text.strip()

        if raw_output.startswith("```json"):
            raw_output = raw_output[len("```json"):].strip()

        elif raw_output.startswith("```"):
            raw_output = raw_output[len("```"):].strip()

        if raw_output.endswith("```"):
            raw_output = raw_output[:-3].strip()

        start = raw_output.find("{")
        end = raw_output.rfind("}")

        if start != -1 and end != -1 and end > start:
            raw_output = raw_output[start:end + 1]

        result = json.loads(raw_output)

        question = result.get("question")

        if not isinstance(question, str) or not question.strip():
            raise HTTPException(
                status_code=500,
                detail="AI returned an empty question.",
            )

        return {
            "question": question.strip()
        }

    except json.JSONDecodeError:
        raise HTTPException(
            status_code=500,
            detail="AI returned invalid JSON.",
        )

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"ExamWarrior AI error: {str(e)}",
        )


# ============================================================
# COPYCATCHER AI
# ============================================================


@app.post("/copycatcher/analyze")
async def copycatcher_analyze(
    file: UploadFile = File(...)
):

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected.",
        )

    filename = file.filename.lower()

    if not filename.endswith((".pdf", ".docx")):
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX files are supported.",
        )

    try:
        file_bytes = await file.read()

        if not file_bytes:
            raise HTTPException(
                status_code=400,
                detail="Uploaded file is empty.",
            )

        if filename.endswith(".pdf"):

            from pypdf import PdfReader

            reader = PdfReader(BytesIO(file_bytes))

            text_parts = []

            for page in reader.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)

            document_text = "\n".join(text_parts)

        else:

            from docx import Document

            document = Document(BytesIO(file_bytes))

            text_parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            document_text = "\n".join(text_parts)

        if not document_text.strip():
            raise HTTPException(
                status_code=400,
                detail="Could not extract text from the document.",
            )

        document_text = document_text[:20000]

        prompt = f"""
You are an AI academic originality assistant for a college app.

Analyze the following student assignment/document.

DOCUMENT TEXT:

{document_text}

Estimate:

1. Similarity with commonly available or repeated academic content.
2. Originality of the writing.
3. Probability that the writing appears AI-generated.

Important:

- These are estimates, not definitive proof.
- Do not claim that a document is definitely AI-generated.
- Give practical recommendations.
- Return ONLY valid JSON.

Use exactly this structure:

{{
  "similarity_percentage": 0,
  "originality_percentage": 0,
  "ai_percentage": 0,
  "summary": "",
  "matched_areas": [],
  "recommendations": []
}}

Rules:

- All percentages must be between 0 and 100.
- similarity_percentage + originality_percentage should equal 100.
- matched_areas must be an array of strings.
- recommendations must be an array of strings.
"""

        response = client.responses.create(
            model="gpt-5.6-luna",
            input=prompt,
        )

        result = json.loads(response.output_text)

        return {
            "similarity_percentage": result.get(
                "similarity_percentage",
                0,
            ),
            "originality_percentage": result.get(
                "originality_percentage",
                0,
            ),
            "ai_percentage": result.get(
                "ai_percentage",
                0,
            ),
            "summary": result.get(
                "summary",
                "",
            ),
            "matched_areas": result.get(
                "matched_areas",
                [],
            ),
            "recommendations": result.get(
                "recommendations",
                [],
            ),
        }

    except HTTPException:
        raise

    except json.JSONDecodeError:
        raise HTTPException(
            status_code=500,
            detail="AI returned invalid JSON.",
        )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"CopyCatcher AI error: {str(e)}",
        )# ============================================================
# CAMPUSSHIELD AI
# ============================================================


@app.post("/safety/analyze")
def safety_analyze(request: SafetyAIRequest):

    reports_text = json.dumps(
        request.recent_reports[:50],
        ensure_ascii=False,
    )

    prompt = f"""
You are CampusShield AI, an AI safety-analysis assistant
for a college campus.

Analyze this campus safety report.

Report type:
{request.report_type}

Description:
{request.description}

Location:
Latitude: {request.latitude}
Longitude: {request.longitude}

Recent reports:
{reports_text}

Return ONLY valid JSON:

{{
  "risk_score": 0,
  "risk_level": "Low",
  "category": "",
  "severity": "",
  "reason": "",
  "recommended_action": "",
  "time_risk": ""
}}

Rules:

- risk_score must be between 0 and 100.
- risk_level must be Low, Medium, High, or Critical.
- severity must be Low, Medium, High, or Critical.
- category should describe the main safety issue.
- reason should briefly explain the estimated risk.
- recommended_action should give a practical action for
  campus staff.
- time_risk should mention a time-related concern only when
  supported by the provided information.
- Do not claim certainty.
- This is a risk estimate, not a guarantee of safety.
"""

    try:
        response = client.responses.create(
            model="gpt-5.6-luna",
            input=prompt,
        )

        result = json.loads(response.output_text)

        return {
            "risk_score": result.get(
                "risk_score",
                0,
            ),
            "risk_level": result.get(
                "risk_level",
                "Low",
            ),
            "category": result.get(
                "category",
                "",
            ),
            "severity": result.get(
                "severity",
                "Low",
            ),
            "reason": result.get(
                "reason",
                "",
            ),
            "recommended_action": result.get(
                "recommended_action",
                "",
            ),
            "time_risk": result.get(
                "time_risk",
                "",
            ),
        }

    except json.JSONDecodeError:
        raise HTTPException(
            status_code=500,
            detail="Safety AI returned invalid JSON.",
        )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Safety AI error: {str(e)}",
        )


# ============================================================
# STUDENT ATTENDANCE AI
# ============================================================


@app.post("/attendance/student-insight")
def attendance_student_insight(
    request: AttendanceAIRequest
):

    if request.total_classes <= 0:
        raise HTTPException(
            status_code=400,
            detail="Total classes must be greater than zero.",
        )

    attendance_percentage = (
        request.present_count / request.total_classes
    ) * 100

    records_text = json.dumps(
        request.recent_records[:50],
        ensure_ascii=False,
    )

    prompt = f"""
You are CampusX Attendance AI.

Analyze a student's attendance data and provide a simple
academic attendance insight.

Student:
{request.student_name}

Present classes:
{request.present_count}

Absent classes:
{request.absent_count}

Total classes:
{request.total_classes}

Current attendance percentage:
{attendance_percentage:.1f}%

Recent attendance records:
{records_text}

Return ONLY valid JSON:

{{
  "attendance_percentage": 0,
  "status": "Good",
  "risk_level": "Low",
  "insight": "",
  "recommendation": "",
  "trend": ""
}}

Rules:

- attendance_percentage must be between 0 and 100.
- status must be Good, Attention, or Critical.
- risk_level must be Low, Medium, or High.
- Give practical and non-alarming advice.
- Do not invent attendance records.
"""

    try:
        response = client.responses.create(
            model="gpt-5.6-luna",
            input=prompt,
        )

        result = json.loads(response.output_text)

        return {
            "attendance_percentage": result.get(
                "attendance_percentage",
                round(attendance_percentage, 1),
            ),
            "status": result.get(
                "status",
                "Good",
            ),
            "risk_level": result.get(
                "risk_level",
                "Low",
            ),
            "insight": result.get(
                "insight",
                "",
            ),
            "recommendation": result.get(
                "recommendation",
                "",
            ),
            "trend": result.get(
                "trend",
                "",
            ),
        }

    except json.JSONDecodeError:
        raise HTTPException(
            status_code=500,
            detail="Attendance AI returned invalid JSON.",
        )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Attendance AI error: {str(e)}",
        )


# ============================================================
# TEACHER ATTENDANCE AI
# ============================================================


@app.post("/attendance/teacher-analytics")
def attendance_teacher_analytics(
    request: TeacherAttendanceAIRequest
):

    students_text = json.dumps(
        request.students[:100],
        ensure_ascii=False,
    )

    summary_text = json.dumps(
        request.recent_summary[:50],
        ensure_ascii=False,
    )

    prompt = f"""
You are CampusX Attendance AI for teachers.

Analyze the attendance of a college class.

Class:
{request.class_name}

Total students:
{request.total_students}

Student attendance data:
{students_text}

Recent class summary:
{summary_text}

Return ONLY valid JSON:

{{
  "class_status": "Good",
  "overall_percentage": 0,
  "low_attendance_students": [],
  "high_attendance_students": [],
  "insight": "",
  "recommendation": ""
}}

Rules:

- overall_percentage must be between 0 and 100.
- class_status must be Good, Attention, or Critical.
- low_attendance_students must contain student names or
  identifiers supplied in the input.
- Do not invent students.
- Give a concise teacher-friendly summary.
- Recommendations should focus on attendance improvement.
"""

    try:
        response = client.responses.create(
            model="gpt-5.6-luna",
            input=prompt,
        )

        result = json.loads(response.output_text)

        return {
            "class_status": result.get(
                "class_status",
                "Good",
            ),
            "overall_percentage": result.get(
                "overall_percentage",
                0,
            ),
            "low_attendance_students": result.get(
                "low_attendance_students",
                [],
            ),
            "high_attendance_students": result.get(
                "high_attendance_students",
                [],
            ),
            "insight": result.get(
                "insight",
                "",
            ),
            "recommendation": result.get(
                "recommendation",
                "",
            ),
        }

    except json.JSONDecodeError:
        raise HTTPException(
            status_code=500,
            detail="Teacher Attendance AI returned invalid JSON.",
        )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Teacher Attendance AI error: {str(e)}",
        )# ==============================
# CAMPUSX AI VOICE ASSISTANT
# ==============================

class AssistantRequest(BaseModel):
    role: str
    question: str
    context: dict = {}


@app.post("/assistant/chat")
def assistant_chat(request: AssistantRequest):
    try:
        role = request.role.strip().lower()
        question = request.question.strip()
        context = request.context or {}

        if not question:
            raise HTTPException(
                status_code=400,
                detail="Question is required."
            )

        system_prompt = """
You are CampusX AI, the intelligent voice assistant inside a college
smart-campus application.

You help students, teachers and administrators.

IMPORTANT RULES:
- Answer only using the information available in the provided context.
- Never invent attendance numbers, safety reports, students or other data.
- If the context does not contain the requested information, clearly say
  that the information is currently unavailable.
- Keep answers short because the answer will be spoken aloud.
- Normally use 1 to 3 short sentences.
- Understand English, Hindi and Hinglish questions.
- Reply in the same language style as the user's question.
- For Hindi questions, return natural Hindi in Devanagari script.
- For English questions, return English.
- Do not mention APIs, databases, JSON, backend or internal implementation.
- Be friendly and helpful.

ROLE:
The user can be a student, teacher or admin.

CONTEXT:
The Flutter app provides trusted campus data in the context object.
Use that data when answering.
"""

        user_prompt = f"""
User role:
{role}

User question:
{question}

Available CampusX context:
{json.dumps(context, ensure_ascii=False, default=str)}
"""

        response = client.chat.completions.create(
            model="gpt-5.6-luna",
            messages=[
                {
                    "role": "system",
                    "content": system_prompt,
                },
                {
                    "role": "user",
                    "content": user_prompt,
                },
            ],
            temperature=0.2,
        )

        reply = response.choices[0].message.content

        if not reply:
            raise HTTPException(
                status_code=500,
                detail="AI returned an empty response."
            )

        reply = reply.strip()

        # Detect response language for Flutter TTS.
        hindi_chars = any(
            "\u0900" <= char <= "\u097F"
            for char in reply
        )

        language = "hi" if hindi_chars else "en"

        return {
            "reply": reply,
            "language": language,
            "role": role,
        }

    except HTTPException:
        raise

    except Exception as e:
        print(f"Assistant error: {e}")

        raise HTTPException(
            status_code=500,
            detail="CampusX AI assistant is temporarily unavailable."
        )